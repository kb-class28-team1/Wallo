from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document


def paragraph_payload(paragraph):
    return {
        "style": paragraph.style.name if paragraph.style is not None else None,
        "text": paragraph.text,
    }


def table_payload(table):
    rows = []
    for row in table.rows:
        rows.append([cell.text for cell in row.cells])
    return rows


def main() -> None:
    input_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])
    doc = Document(str(input_path))
    payload = {
        "core_properties": {
            "title": doc.core_properties.title,
            "subject": doc.core_properties.subject,
            "author": doc.core_properties.author,
            "last_modified_by": doc.core_properties.last_modified_by,
        },
        "paragraphs": [paragraph_payload(p) for p in doc.paragraphs],
        "tables": [table_payload(t) for t in doc.tables],
        "sections": len(doc.sections),
        "inline_shapes": len(doc.inline_shapes),
    }
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
